from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


def create_system_architecture_doc():
    # Khởi tạo tài liệu Word
    doc = Document()

    # --- TIÊU ĐỀ CHÍNH ---
    title = doc.add_heading('BẢN THIẾT KẾ KIẾN TRÚC HỆ THỐNG', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph(
        'Dự án: Hệ thống Nhận diện Ngôn ngữ Ký hiệu Tĩnh (HCI)\n').alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # --- PHẦN I: LUỒNG HOẠT ĐỘNG ---
    doc.add_heading('PHẦN I: GIAI ĐOẠN HUẤN LUYỆN (OFFLINE TRAINING PIPELINE)', level=1)
    doc.add_paragraph('Mục tiêu: Xử lý dữ liệu tĩnh và tạo ra các file mô hình (.pkl) để lưu trữ.')

    steps_phase1 = [
        ("Bước 1: Thu thập Dữ liệu Thô (Raw Data Collection)",
         "Quét toàn bộ thư mục ảnh tĩnh (image_dataset). Sử dụng MediaPipe Tasks API (chế độ IMAGE) trích xuất tọa độ 21 điểm mốc của bàn tay thuận. Lưu kết quả ra file dataset_raw.csv."),
        ("Bước 2: Tiền xử lý & Chuẩn hóa (Data Preprocessing)",
         "Đọc dữ liệu thô. Áp dụng toán học không gian: Dời tâm tọa độ về gốc (0,0,0) tại cổ tay và Thu phóng (Scaling) tỷ lệ khung xương. Xuất ra file dataset_normalized.csv."),
        ("Bước 3: Huấn luyện Mô hình (Model Training)",
         "Sử dụng dữ liệu đã chuẩn hóa để huấn luyện thuật toán Random Forest hoặc SVM (kèm StandardScaler để tối ưu không gian phân bố)."),
        ("Bước 4: Đóng gói và Xuất xưởng (Model Exporting)",
         "Dùng thư viện joblib xuất các trọng số đã học thành file nhị phân (rf_model.pkl, svm_model.pkl, scaler.pkl) vào thư mục exported_models/.")
    ]
    for step_title, step_desc in steps_phase1:
        p = doc.add_paragraph()
        p.add_run(step_title).bold = True
        p.add_run(f'\n{step_desc}')

    doc.add_heading('PHẦN II: GIAI ĐOẠN VẬN HÀNH (REAL-TIME WEB SERVICE)', level=1)
    doc.add_paragraph('Mục tiêu: Đưa hệ thống lên trình duyệt, phân tách giao diện và logic suy luận AI.')

    steps_phase2 = [
        ("Bước 5: Khởi tạo Giao diện & Bắt chuyển động (Frontend)",
         "Trình duyệt xin quyền Webcam, tải file hand_landmarker.task và chạy MediaPipe cục bộ bằng phần cứng người dùng. Cấu hình nhận diện tối đa 2 bàn tay."),
        ("Bước 6: Xử lý Logic Tương tác Hai tay (HCI Logic)",
         "Bàn tay phụ làm nút Enter (Tính khoảng cách ngón tay để xác định trạng thái NẮM/MỞ). Bàn tay thuận làm ký hiệu, chỉ đóng gói tọa độ thô khi tay phụ MỞ."),
        ("Bước 7: Giao tiếp Hệ thống (Client-Server Request)",
         "Frontend bắn gói JSON chứa mảng tọa độ qua HTTP POST/WebSocket lên FastAPI Server. Tuyệt đối không gửi hình ảnh video."),
        ("Bước 8: Xử lý Trực tiếp tại Máy chủ (Backend Preprocessing)",
         "FastAPI nhận mảng JSON. Gọi hàm chuẩn hóa (Dời tâm & Thu phóng) từ utils.py để xử lý tọa độ thô thành dữ liệu sạch ngay trên RAM."),
        ("Bước 9: Suy luận & Trả kết quả (Inference & Response)",
         "Đưa mảng vừa chuẩn hóa vào mô hình đã load (model.predict). Đóng gói nhãn dự đoán trả về Frontend hiển thị.")
    ]
    for step_title, step_desc in steps_phase2:
        p = doc.add_paragraph()
        p.add_run(step_title).bold = True
        p.add_run(f'\n{step_desc}')

    doc.add_page_break()

    # --- PHẦN III: CẤU TRÚC THƯ MỤC ---
    doc.add_heading('PHẦN III: CẤU TRÚC THƯ MỤC DỰ ÁN (MÔI TRƯỜNG PYCHARM)', level=1)

    directory_structure = """SignLanguage_System/
├── venv/                              <-- Thư mục môi trường ảo
│
├── Phase1_OfflineTraining/            <-- [HUẤN LUYỆN]
│   ├── image_dataset/                 <-- Thư mục chứa ảnh gốc (A, B, C...)
│   ├── data/                          <-- Chứa dataset_raw.csv & dataset_normalized.csv
│   ├── exported_models/               <-- Chứa rf_model.pkl, svm_model.pkl, scaler.pkl
│   │
│   ├── 1_extract_landmarks.py         <-- Quét ảnh, lưu CSV thô
│   ├── 2_preprocess_data.py           <-- Chuẩn hóa dữ liệu thô thành sạch
│   ├── 3a_train_rf.py                 <-- Train Random Forest
│   ├── 3b_train_svm.py                <-- Train SVM
│   ├── utils.py                       <-- Các hàm toán học chuẩn hóa (Dời tâm, Thu phóng)
│   └── hand_landmarker.task           <-- Model MediaPipe cục bộ
│
└── Phase2_OnlineWebService/           <-- [VẬN HÀNH WEB API]
    ├── frontend/                      
    │   ├── index.html                 <-- Giao diện UI
    │   ├── app.js                     <-- Xử lý luồng 2 tay, gọi API
    │   └── hand_landmarker.task       <-- Model MediaPipe cho Web
    │
    └── backend/                       
        ├── main.py                    <-- API Server (FastAPI)
        ├── utils.py                   <-- Copy từ Phase 1 sang để chuẩn hóa Real-time
        └── models/                    <-- Copy từ Phase 1 sang để load vào RAM
"""

    # Định dạng font chữ dạng Code cho cấu trúc thư mục dễ nhìn
    p_code = doc.add_paragraph()
    run = p_code.add_run(directory_structure)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)

    # --- LƯU FILE ---
    filename = "Cau_Truc_He_Thong_Nhan_Dien.docx"
    doc.save(filename)
    print(f"[+] Tạo thành công tệp Word: {filename}")


if __name__ == "__main__":
    create_system_architecture_doc()